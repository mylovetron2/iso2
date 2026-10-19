"""Phần mềm tổng hợp KPI SC/BD/CC thiết bị.

Dán dữ liệu copy từ Excel (các khối "Nhân viên: ..."), bấm "In" để xuất file Excel
gộp theo số phiếu (mẫu KPI - mới).
"""
import calendar
import csv
import io
import os
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = [
    "Phiếu", "Hồ sơ", "Mã VT", "Số máy", "Ngày YC", "Ngày TH", "Ngày KT",
    "Số giờ định mức", "Số giờ thực hiện", "KPI", "Kết quả",
]
WIDTHS = [9, 60, 22, 10, 12, 12, 12, 11, 11, 14, 22]
NCOL = len(HEADERS)
SPECIAL_NAMES = {"VŨ ANH ĐỨC": "A.Đức", "ĐOÀN MINH ĐỨC": "M.Đức"}
NO_KPI = "chưa gán KPI"

FONT = Font(name="Times New Roman", size=11)
BOLD = Font(name="Times New Roman", size=11, bold=True)
THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL = PatternFill("solid", fgColor="87CEEB")


# ---------------------------------------------------------------- phân tích dữ liệu
def parse_date(text):
    text = (text or "").strip()
    if not text:
        return None
    for fmt in ("%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def parse_number(text):
    text = (text or "").strip().replace(",", ".")
    if not text:
        return None
    try:
        value = float(text)
    except ValueError:
        return None
    return int(value) if value.is_integer() else value


def clean(text):
    return re.sub(r"\s+", " ", (text or "")).strip()


def parse_input(text):
    """Trả về (records, warnings); record = dict gồm nhân viên + 11 cột dữ liệu."""
    records, warnings = [], []
    employee = ""
    reader = csv.reader(io.StringIO(text.replace("\r\n", "\n").replace("\r", "\n")), delimiter="\t")
    for cells in reader:
        cells = [clean(c) for c in cells]
        if not any(cells):
            continue
        first = cells[0]
        match = re.match(r"^nhân\s*viên\s*:\s*(.*)$", first, re.IGNORECASE)
        if match:
            employee = match.group(1).strip()
            continue
        if first.lower().startswith("phiếu"):
            continue
        if len(cells) < NCOL:
            cells += [""] * (NCOL - len(cells))
        if not first:
            warnings.append("Bỏ qua dòng không có số phiếu: " + " | ".join(c for c in cells if c)[:80])
            continue
        if len(cells) > NCOL:
            warnings.append(f"Phiếu {first}: dư cột, các cột sau cột Kết quả bị bỏ qua.")
        phieu, hoso, mavt, somay, yc, th, kt, dm, gio, kpi, kq = cells[:NCOL]
        for label, raw in (("Ngày YC", yc), ("Ngày TH", th), ("Ngày KT", kt)):
            if raw and parse_date(raw) is None:
                warnings.append(f"Phiếu {first}: {label} '{raw}' không đọc được (giữ nguyên chữ).")
        for label, raw in (("Số giờ định mức", dm), ("Số giờ thực hiện", gio)):
            if raw and parse_number(raw) is None:
                warnings.append(f"Phiếu {first}: {label} '{raw}' không phải số.")
        records.append({
            "nv": employee, "phieu": phieu, "hoso": hoso, "mavt": mavt, "somay": somay,
            "yc": yc, "th": th, "kt": kt, "dm": parse_number(dm), "gio": parse_number(gio),
            "kpi": kpi, "kq": kq,
        })
    return records, warnings


def phieu_key(phieu):
    parts = re.split(r"[.\-/]", phieu)
    return tuple((0, int(p), "") if p.isdigit() else (1, 0, p) for p in parts)


def short_names(employees):
    """Tên ngắn như hệ thống (từ cuối); trùng tên thì thêm chữ cái đầu của chữ lót."""
    employees = list(dict.fromkeys(e for e in employees if e))
    counts = {}
    for e in employees:
        if e not in SPECIAL_NAMES:
            last = e.split()[-1].title()
            counts[last] = counts.get(last, 0) + 1
    result = {}
    for e in employees:
        if e in SPECIAL_NAMES:
            result[e] = SPECIAL_NAMES[e]
            continue
        words = e.split()
        last = words[-1].title()
        if len(words) > 2 and counts[last] > 1:
            result[e] = words[-2][0].upper() + "." + last
        else:
            result[e] = last
    return result


def merge_by_phieu(records, use_short_names=True):
    names = short_names(r["nv"] for r in records) if use_short_names else {}
    groups = {}
    for rec in records:
        groups.setdefault(rec["phieu"], []).append(rec)

    merged = []
    for phieu in sorted(groups, key=phieu_key):
        rows = groups[phieu]
        item = {"phieu": phieu}
        for field in ("hoso", "mavt", "somay", "yc", "th", "kt"):
            item[field] = next((r[field] for r in rows if r[field]), "")
        employees = list(dict.fromkeys(r["nv"] for r in rows if r["nv"]))
        item["nv"] = ", ".join(dict.fromkeys(names.get(e, e) for e in employees))
        dms = [r["dm"] for r in rows if r["dm"] is not None]
        gios = [r["gio"] for r in rows if r["gio"] is not None]
        item["dm"] = max(dms) if dms else None
        item["gio"] = max(gios) if gios else None
        kpis = {r["kpi"] for r in rows if r["kpi"]}
        if NO_KPI in kpis:
            item["kpi"] = NO_KPI
        elif len(kpis) == 1:
            item["kpi"] = next(iter(kpis))
        else:
            item["kpi"] = "/".join(sorted(kpis))
        item["kq"] = next((r["kq"] for r in rows if r["kq"]), "")
        merged.append(item)
    return merged


def guess_period(records):
    """Mặc định: cả tháng của ngày KT lớn nhất (không có thì ngày YC lớn nhất)."""
    for field in ("kt", "yc"):
        dates = [d for d in (parse_date(r[field]) for r in records) if d]
        if dates:
            last = max(dates)
            return (datetime(last.year, last.month, 1),
                    datetime(last.year, last.month, calendar.monthrange(last.year, last.month)[1]))
    return None, None


# ---------------------------------------------------------------- xuất Excel
def date_value(text):
    return parse_date(text) or (text or None)


def write_title(ws, last_col, title, period_text):
    ws["A1"] = "XN ĐỊA VẬT LÝ GK"
    ws["A1"].font = BOLD
    ws["A2"] = "XƯỞNG SCTBĐVL"
    ws["A2"].font = FONT
    for row in (1, 2):
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=last_col)
    ws["C1"] = title
    ws["C1"].font = BOLD
    ws["C2"] = period_text
    ws["C2"].font = FONT
    for cell in ("C1", "C2"):
        ws[cell].alignment = Alignment(horizontal="center")


def write_header(ws, row, headers):
    for col, text in enumerate(headers, 1):
        cell = ws.cell(row, col, text)
        cell.font = BOLD
        cell.fill = FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_row(ws, row, values, date_cols, center_cols):
    for col, value in enumerate(values, 1):
        cell = ws.cell(row, col, value)
        cell.font = FONT
        cell.border = BORDER
        if col in date_cols:
            cell.number_format = "dd/mm/yyyy"
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col in center_cols:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        else:
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def setup_sheet(ws, widths, header_row):
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if header_row:
        ws.print_title_rows = f"{header_row}:{header_row}"


def build_workbook(records, start, end, use_short_names=True, with_detail=True):
    period_text = f"Từ {start:%d/%m/%Y} đến {end:%d/%m/%Y}"
    wb = Workbook()

    ms = wb.active
    ms.title = "Gộp theo phiếu"
    merged_headers = HEADERS[:7] + ["Nhân viên thực hiện"] + HEADERS[7:]
    merged_widths = WIDTHS[:7] + [30] + WIDTHS[7:]
    write_title(ms, len(merged_headers),
                "TỔNG HỢP KPI CÔNG TÁC BẢO DƯỠNG, SỬA CHỮA, CHUẨN CHỈNH THIẾT BỊ (GỘP THEO PHIẾU)",
                period_text)
    write_header(ms, 4, merged_headers)
    merged = merge_by_phieu(records, use_short_names)
    for i, m in enumerate(merged, 5):
        write_row(ms, i, [
            m["phieu"], m["hoso"], m["mavt"], m["somay"],
            date_value(m["yc"]), date_value(m["th"]), date_value(m["kt"]),
            m["nv"], m["dm"], m["gio"], m["kpi"], m["kq"],
        ], date_cols=(5, 6, 7), center_cols=(1, 4, 9, 10, 11))
    setup_sheet(ms, merged_widths, 4)

    if with_detail:
        ws = wb.create_sheet("KPI theo nhân viên")
        write_title(ws, NCOL, "TỔNG HỢP KPI CÔNG TÁC BẢO DƯỠNG, SỬA CHỮA, CHUẨN CHỈNH THIẾT BỊ THEO NHÂN VIÊN",
                    period_text)
        row = 4
        by_employee = {}
        for rec in records:
            by_employee.setdefault(rec["nv"], []).append(rec)
        for index, (employee, items) in enumerate(by_employee.items()):
            if index:
                row += 1
            ws.cell(row, 1, "Nhân viên: " + employee).font = BOLD
            row += 1
            write_header(ws, row, HEADERS)
            row += 1
            for rec in items:
                write_row(ws, row, [
                    rec["phieu"], rec["hoso"], rec["mavt"], rec["somay"],
                    date_value(rec["yc"]), date_value(rec["th"]), date_value(rec["kt"]),
                    rec["dm"], rec["gio"], rec["kpi"], rec["kq"],
                ], date_cols=(5, 6, 7), center_cols=(1, 4, 8, 9, 10))
                row += 1
        setup_sheet(ws, WIDTHS, None)
    return wb, len(merged)


# ---------------------------------------------------------------- giao diện
def run_gui():
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    root = tk.Tk()
    root.title("Tổng hợp KPI SC/BD/CC thiết bị")
    root.geometry("1000x640")

    top = ttk.Frame(root, padding=8)
    top.pack(fill="x")
    ttk.Label(top, text="Dán dữ liệu copy từ Excel (Ctrl+V) vào ô bên dưới, rồi bấm In.").pack(side="left")

    text = tk.Text(root, wrap="none", undo=True, font=("Consolas", 10))
    ysb = ttk.Scrollbar(root, orient="vertical", command=text.yview)
    xsb = ttk.Scrollbar(root, orient="horizontal", command=text.xview)
    text.configure(yscrollcommand=ysb.set, xscrollcommand=xsb.set)
    xsb.pack(side="bottom", fill="x")
    ysb.pack(side="right", fill="y")

    opts = ttk.Frame(root, padding=8)
    opts.pack(side="bottom", fill="x")
    text.pack(fill="both", expand=True, padx=8)

    from_var, to_var = tk.StringVar(), tk.StringVar()
    short_var, detail_var = tk.BooleanVar(value=True), tk.BooleanVar(value=True)
    status = tk.StringVar(value="Chưa có dữ liệu.")

    ttk.Label(opts, text="Từ ngày:").grid(row=0, column=0, sticky="w")
    ttk.Entry(opts, textvariable=from_var, width=12).grid(row=0, column=1, padx=4)
    ttk.Label(opts, text="Đến ngày:").grid(row=0, column=2, sticky="w", padx=(12, 0))
    ttk.Entry(opts, textvariable=to_var, width=12).grid(row=0, column=3, padx=4)
    ttk.Label(opts, text="(dd/mm/yyyy, để trống = tự lấy cả tháng của Ngày KT)").grid(
        row=0, column=4, sticky="w", padx=8)
    ttk.Checkbutton(opts, text="Tên nhân viên rút gọn", variable=short_var).grid(
        row=1, column=0, columnspan=3, sticky="w", pady=4)
    ttk.Checkbutton(opts, text="Kèm sheet chi tiết theo nhân viên", variable=detail_var).grid(
        row=1, column=3, columnspan=3, sticky="w", pady=4)
    ttk.Label(opts, textvariable=status).grid(row=2, column=0, columnspan=5, sticky="w")

    def paste():
        try:
            data = root.clipboard_get()
        except tk.TclError:
            messagebox.showwarning("Dán", "Clipboard đang trống.")
            return
        text.delete("1.0", "end")
        text.insert("1.0", data)
        refresh_status()

    def clear():
        text.delete("1.0", "end")
        status.set("Chưa có dữ liệu.")

    def refresh_status(_event=None):
        records, _ = parse_input(text.get("1.0", "end"))
        if records:
            phieus = len({r["phieu"] for r in records})
            status.set(f"Đọc được {len(records)} dòng, {len({r['nv'] for r in records})} nhân viên, {phieus} phiếu.")
        else:
            status.set("Chưa có dữ liệu.")

    def parse_period(records):
        start, end = parse_date(from_var.get()), parse_date(to_var.get())
        if from_var.get().strip() and not start or to_var.get().strip() and not end:
            raise ValueError("Ngày không hợp lệ, nhập theo dd/mm/yyyy.")
        if not start or not end:
            gs, ge = guess_period(records)
            start, end = start or gs, end or ge
        if not start or not end:
            raise ValueError("Không xác định được khoảng ngày, hãy nhập Từ ngày / Đến ngày.")
        if start > end:
            start, end = end, start
        return start, end

    def do_print():
        records, warnings = parse_input(text.get("1.0", "end"))
        if not records:
            messagebox.showwarning("In", "Chưa có dữ liệu hợp lệ để in.")
            return
        try:
            start, end = parse_period(records)
        except ValueError as exc:
            messagebox.showwarning("In", str(exc))
            return
        path = filedialog.asksaveasfilename(
            title="Lưu file Excel", defaultextension=".xlsx",
            initialfile=f"KPI_tonghop_{end:%m-%Y}.xlsx", filetypes=[("Excel", "*.xlsx")])
        if not path:
            return
        try:
            wb, count = build_workbook(records, start, end, short_var.get(), detail_var.get())
            wb.save(path)
        except PermissionError:
            messagebox.showerror("In", "Không ghi được file. Hãy đóng file Excel đang mở rồi thử lại.")
            return
        except Exception as exc:  # noqa: BLE001
            messagebox.showerror("In", f"Lỗi khi tạo file: {exc}")
            return
        message = f"Đã xuất {count} phiếu:\n{path}"
        if warnings:
            message += "\n\nCảnh báo:\n" + "\n".join(warnings[:15])
            if len(warnings) > 15:
                message += f"\n... và {len(warnings) - 15} cảnh báo khác"
        messagebox.showinfo("In", message)
        try:
            os.startfile(path)  # noqa: S606
        except (AttributeError, OSError):
            pass

    ttk.Button(top, text="Xóa", command=clear).pack(side="right", padx=4)
    ttk.Button(top, text="Dán", command=paste).pack(side="right", padx=4)
    ttk.Button(top, text="In (xuất Excel)", command=do_print).pack(side="right", padx=4)
    text.bind("<<Paste>>", lambda e: root.after(50, refresh_status))
    text.bind("<KeyRelease>", refresh_status)

    root.mainloop()


if __name__ == "__main__":
    if len(sys.argv) == 3:
        # Chạy dòng lệnh: kpi_tonghop.py <file_du_lieu.txt> <file_ra.xlsx>
        with open(sys.argv[1], encoding="utf-8-sig") as fh:
            recs, warns = parse_input(fh.read())
        s, e = guess_period(recs)
        workbook, n = build_workbook(recs, s, e)
        workbook.save(sys.argv[2])
        print(f"{len(recs)} dòng -> {n} phiếu; cảnh báo: {len(warns)}")
    else:
        run_gui()
